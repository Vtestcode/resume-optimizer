"""
Resume Aligner — Streamlit Application  (v4)
=============================================
Aligns a Word (.docx) resume with a job description using the OpenAI API.

CORE DESIGN:
  Many resumes pack entire sections into a SINGLE Word paragraph using
  <w:br/> (line-break) elements.  Each \\n in paragraph.text maps to a <w:br/>.

  Strategy (v4 — formatting-safe):
    1.  Split each paragraph into "run-groups" separated by <w:br/>.
        Each run-group = one visual line.
    2.  Classify every line (heading / role_title / bullet / skill_line / summary_bullet / text).
    3.  Send bullets & summary bullets to OpenAI for rewriting per role / section.
    4.  For CHANGED lines only: rebuild that line's run-group from new text.
        For UNCHANGED lines: deep-copy the original run elements byte-for-byte.
    5.  Reassemble the paragraph with original <w:br/> separators intact,
        preserving <w:pPr> (paragraph properties) and per-line formatting.

  What changed from v3:
    • Each line keeps its OWN base formatting (sz=28 for headings, default for bullets).
    • Unchanged lines are cloned verbatim — no formatting loss.
    • Professional Summary section bullets are now aligned too.
    • Skills section preserves bold category names.

Tech: Python · Streamlit · OpenAI · python-docx · pandas · dotenv
"""

# ──────────────────────────────────────────────────────────────────────
# IMPORTS
# ──────────────────────────────────────────────────────────────────────
import os, io, re, json, copy, math
from typing import List, Dict, Tuple, Optional, Any
from collections import OrderedDict

import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from dotenv import load_dotenv
from openai import OpenAI
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

load_dotenv()
DEFAULT_MODEL = "gpt-4o"
METRIC_PATTERN = re.compile(
    r"(\b\d+(?:\.\d+)?%|\b\d+(?:\.\d+)?(?:[kKmMbB]|x)\b|\b\d[\d,]*(?:\.\d+)?\b|\bSLA\b|\bSLI\b|\bSLO\b|\bms\b|\bsec\b|\bsecs\b|\bseconds\b|\bminutes\b|\bhours\b|\bdays\b|\$\s?\d)",
    re.IGNORECASE,
)
STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "in", "into",
    "of", "on", "or", "that", "the", "to", "with", "using", "used", "built",
    "managed", "delivered", "improved", "supported", "led", "across", "over",
}


# ======================================================================
# LOW-LEVEL XML HELPERS
# ======================================================================

def _collect_run_groups(para) -> List[List]:
    """
    Split a paragraph's runs into visual lines.

    Word paragraphs may encode a new line either as a real <w:br/> element
    or as a literal "\\n" inside <w:t>.  The parser classifies lines using
    paragraph.text, so the rebuilder must split lines the same way or line
    indexes drift and title/client lines can disappear during rewrites.
    """
    groups: List[List] = [[]]
    for r in para._element.findall(qn("w:r")):
        segments = _split_run_segments(r)
        if not segments:
            continue
        groups[-1].extend(segments[0])
        for segment in segments[1:]:
            groups.append([])
            groups[-1].extend(segment)
    return groups


def _clone_run_with_children(orig_run, children: List):
    """Clone a run shell and attach the supplied child elements."""
    if not children:
        return None
    cloned = etree.Element(qn("w:r"))
    rPr = orig_run.find(qn("w:rPr"))
    if rPr is not None:
        cloned.append(copy.deepcopy(rPr))
    for child in children:
        cloned.append(child)
    return cloned


def _split_run_segments(orig_run) -> List[List]:
    """
    Return line segments for a single run.

    Each item in the returned list represents one visual line contribution
    from the run. Empty items are valid and represent leading/trailing line
    breaks with no text on that side.
    """
    segments: List[List] = [[]]
    current_children: List = []

    def flush_current():
        cloned_run = _clone_run_with_children(orig_run, current_children)
        if cloned_run is not None:
            segments[-1].append(cloned_run)

    for child in orig_run:
        if child.tag == qn("w:rPr"):
            continue
        if child.tag == qn("w:br"):
            flush_current()
            current_children = []
            segments.append([])
            continue
        if child.tag == qn("w:t"):
            text = child.text or ""
            parts = text.split("\n")
            for idx, part in enumerate(parts):
                if part:
                    t = copy.deepcopy(child)
                    t.text = part
                    current_children.append(t)
                if idx < len(parts) - 1:
                    flush_current()
                    current_children = []
                    segments.append([])
            continue
        current_children.append(copy.deepcopy(child))

    flush_current()
    return segments


def _group_plain_text(group: list) -> str:
    """Concatenate all <w:t> text from a run-group."""
    parts = []
    for r in group:
        for t in r.findall(qn("w:t")):
            parts.append(t.text or "")
    return "".join(parts)


def _get_non_bold_rpr(group: list):
    """Get the <w:rPr> from the first non-bold text run in a group (for bullet body text)."""
    for r in group:
        txt = "".join((t.text or "") for t in r.findall(qn("w:t")))
        if not txt.strip():
            continue
        rPr = r.find(qn("w:rPr"))
        if rPr is not None:
            # Check if this run is bold
            b = rPr.find(qn("w:b"))
            is_bold = (b is not None and b.get(qn("w:val")) not in ("0", "false"))
            if not is_bold:
                return copy.deepcopy(rPr)
    # Fallback: return any rPr (strip bold from it)
    for r in group:
        rPr = r.find(qn("w:rPr"))
        if rPr is not None:
            rPr_copy = copy.deepcopy(rPr)
            for tag in (qn("w:b"), qn("w:bCs")):
                el = rPr_copy.find(tag)
                if el is not None:
                    rPr_copy.remove(el)
            return rPr_copy
    return None


def _get_bold_rpr(group: list):
    """Get the <w:rPr> from the first bold text run in a group."""
    for r in group:
        txt = "".join((t.text or "") for t in r.findall(qn("w:t")))
        if not txt.strip():
            continue
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            continue
        b = rPr.find(qn("w:b"))
        is_bold = (b is not None and b.get(qn("w:val")) not in ("0", "false"))
        if is_bold:
            return copy.deepcopy(rPr)
    return None


def _set_bold(rPr_template, bold: bool):
    """Return a copy of rPr with bold toggled."""
    if rPr_template is None:
        rPr = etree.Element(qn("w:rPr"))
    else:
        rPr = copy.deepcopy(rPr_template)
    for tag in (qn("w:b"), qn("w:bCs")):
        existing = rPr.find(tag)
        if bold:
            if existing is None:
                el = etree.SubElement(rPr, tag)
                el.set(qn("w:val"), "1")
            else:
                existing.set(qn("w:val"), "1")
        else:
            if existing is not None:
                rPr.remove(existing)
    return rPr


def _make_text_run(text: str, rPr):
    """Create a <w:r> with a single <w:t>."""
    r = etree.Element(qn("w:r"))
    if rPr is not None:
        r.append(copy.deepcopy(rPr))
    t = etree.SubElement(r, qn("w:t"))
    t.set(qn("xml:space"), "preserve")
    t.text = text
    return r


def _make_br_run(rPr=None):
    """Create a <w:r> with only <w:br/>."""
    r = etree.Element(qn("w:r"))
    if rPr is not None:
        r.append(copy.deepcopy(rPr))
    etree.SubElement(r, qn("w:br"))
    return r


def _parse_bold_markers(text: str) -> List[Tuple[str, bool]]:
    """
    'Built **ETL** with **Python**'
    → [('Built ', False), ('ETL', True), (' with ', False), ('Python', True)]
    """
    segs = []
    last = 0
    for m in re.finditer(r"\*\*(.*?)\*\*", text):
        if m.start() > last:
            segs.append((text[last:m.start()], False))
        segs.append((m.group(1), True))
        last = m.end()
    if last < len(text):
        segs.append((text[last:], False))
    return segs if segs else [(text, False)]


def _build_replacement_runs(new_text: str, base_rPr, prepend_bullet: bool = True) -> list:
    """
    Build <w:r> elements for a replacement line.
    • Parses **bold** markers.
    • Prepends "• " if prepend_bullet.
    • base_rPr = formatting template from the ORIGINAL line's non-bold runs.
    """
    full = f"• {new_text}" if prepend_bullet else new_text
    segments = _parse_bold_markers(full)
    runs = []
    for seg_text, seg_bold in segments:
        if not seg_text:
            continue
        rPr = _set_bold(base_rPr, seg_bold)
        runs.append(_make_text_run(seg_text, rPr))
    return runs


def _build_skill_replacement_runs(new_text: str, group: list) -> list:
    """
    Rebuild a skills line while preserving the original bold category label.
    """
    label_rPr = _get_bold_rpr(group)
    body_rPr = _get_non_bold_rpr(group)
    if body_rPr is None:
        body_rPr = label_rPr
    if label_rPr is None:
        label_rPr = _set_bold(body_rPr, True)

    label, sep, rest = new_text.partition(":")
    if not sep:
        return _build_replacement_runs(new_text, body_rPr, prepend_bullet=False)

    runs = [_make_text_run(f"{label.strip()}:", label_rPr)]
    remainder = rest.lstrip()
    if remainder:
        runs.extend(_build_replacement_runs(f" {remainder}", body_rPr, prepend_bullet=False))
    return runs


def _clean_skills_output(raw: str) -> List[str]:
    """Remove code fences / wrapper labels from the skills response."""
    cleaned = raw.strip()
    cleaned = re.sub(r"^```[A-Za-z0-9_-]*\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
    while lines and re.fullmatch(r"(plaintext|plain\s*text|text)", lines[0], flags=re.I):
        lines.pop(0)
    return lines


def _normalize_line_text(text: str, is_bullet: bool) -> str:
    """Normalize line text for comparing intended vs actual DOCX content."""
    cleaned = re.sub(r"\*\*", "", text or "").strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    if is_bullet:
        cleaned = cleaned.lstrip("•").strip()
    return cleaned


# ======================================================================
# PARAGRAPH REBUILDER  (v4 — formatting-safe)
# ======================================================================

class ParagraphRebuilder:
    """
    Rebuild a paragraph's XML, replacing only the lines specified in
    `line_replacements` while cloning all other lines verbatim.
    """

    @staticmethod
    def rebuild(para, line_replacements: Dict[int, str],
                is_bullet: bool = True, is_skill: bool = False):
        """
        para: a python-docx Paragraph object.
        line_replacements: { line_index: "new text with **bold** markers" }
        is_bullet: True → prepend "• " to replacement text.
        """
        p_elem = para._element

        # 1. Collect the current run-groups (one per visual line)
        groups = _collect_run_groups(para)

        # 2. Build the new children list
        new_children = []

        # Preserve <w:pPr> if present
        pPr = p_elem.find(qn("w:pPr"))
        if pPr is not None:
            new_children.append(copy.deepcopy(pPr))

        # Get a generic rPr for <w:br/> separator runs (from any non-bold run)
        br_rPr = None
        for g in groups:
            br_rPr = _get_non_bold_rpr(g)
            if br_rPr is not None:
                break

        for g_idx, group in enumerate(groups):
            # Insert a <w:br/> before every line except the first
            if g_idx > 0:
                new_children.append(_make_br_run(br_rPr))

            if g_idx in line_replacements:
                # ── REPLACED LINE ──
                # Use the formatting from THIS line's original non-bold runs
                line_base_rPr = _get_non_bold_rpr(group)
                if line_base_rPr is None:
                    line_base_rPr = br_rPr  # fallback

                new_text = line_replacements[g_idx]
                if is_skill:
                    replacement_runs = _build_skill_replacement_runs(new_text, group)
                else:
                    replacement_runs = _build_replacement_runs(
                        new_text, line_base_rPr, prepend_bullet=is_bullet
                    )
                new_children.extend(replacement_runs)
            else:
                # ── UNCHANGED LINE — clone original runs verbatim ──
                for orig_r in group:
                    cloned = copy.deepcopy(orig_r)
                    # Remove any <w:br/> that leaked into this run
                    # (we handle breaks ourselves as separators above)
                    for br in cloned.findall(qn("w:br")):
                        cloned.remove(br)
                    new_children.append(cloned)

        # 3. Clear the paragraph element and re-populate
        for child in list(p_elem):
            p_elem.remove(child)
        for child in new_children:
            p_elem.append(child)


# ======================================================================
# DOCUMENT PARSER
# ======================================================================

class DocumentParser:
    EXPERIENCE_KW = ["experience", "professional experience", "work experience"]
    SUMMARY_KW = ["professional summary", "summary", "profile", "objective"]
    SKILLS_KW = ["skill", "professional skills", "technical skills",
                 "tech stack", "core competencies", "tools", "technologies"]

    @staticmethod
    def load_document(file_bytes: bytes) -> Document:
        return Document(io.BytesIO(file_bytes))

    @classmethod
    def parse_paragraphs(cls, doc: Document) -> List[Dict]:
        """
        Split each paragraph on \\n into visual lines and classify each.
        """
        results = []
        current_section = None

        for pidx, para in enumerate(doc.paragraphs):
            full_text = para.text
            lines = full_text.split("\n")
            classifications = []

            for line in lines:
                stripped = line.strip()
                lower = stripped.lower()

                if not stripped:
                    classifications.append("text")
                    continue

                # Section heading?
                all_section_kw = cls.EXPERIENCE_KW + cls.SUMMARY_KW + cls.SKILLS_KW + [
                    "education", "certifications", "projects", "achievements", "awards"
                ]
                if lower in all_section_kw:
                    classifications.append("heading")
                    current_section = lower
                    continue

                # Bullet?
                if stripped.startswith("•"):
                    # Determine bullet type based on section
                    if current_section and any(kw in current_section for kw in
                                               ("summary", "profile", "objective")):
                        classifications.append("summary_bullet")
                    else:
                        classifications.append("bullet")
                    continue

                # Role title: contains "|" with date-like patterns
                if re.search(r"\|.*?(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec|\d{4}|present)",
                             lower):
                    classifications.append("role_title")
                    continue

                # In skills section → skill_line
                if current_section and any(kw in current_section for kw in
                                           ("skill", "tech", "competenc", "tools")):
                    classifications.append("skill_line")
                    continue

                classifications.append("text")

            results.append({
                "para_idx": pidx,
                "lines": lines,
                "classifications": classifications,
                "section": current_section,
            })

        return results

    @classmethod
    def extract_roles(cls, parsed: List[Dict]) -> OrderedDict:
        """Extract bullet points grouped by role."""
        roles = OrderedDict()
        current_role = None
        in_experience = False

        for rec in parsed:
            for i, (line, cls_label) in enumerate(zip(rec["lines"], rec["classifications"])):
                if cls_label == "heading":
                    lower = line.strip().lower()
                    in_experience = any(kw in lower for kw in cls.EXPERIENCE_KW)
                    if not in_experience:
                        current_role = None
                    continue
                if not in_experience:
                    continue
                if cls_label == "role_title":
                    current_role = line.strip()[:150]
                    if current_role not in roles:
                        roles[current_role] = {"para_idx": rec["para_idx"], "bullets": []}
                    continue
                if cls_label == "bullet" and current_role:
                    bullet_text = line.strip().lstrip("•").strip()
                    roles[current_role]["bullets"].append({
                        "para_idx": rec["para_idx"],
                        "line_idx": i,
                        "text": bullet_text,
                    })
        return roles

    @classmethod
    def extract_summary_bullets(cls, parsed: List[Dict]) -> List[Dict]:
        """Extract bullets from Professional Summary section."""
        items = []
        in_summary = False
        for rec in parsed:
            for i, (line, cls_label) in enumerate(zip(rec["lines"], rec["classifications"])):
                if cls_label == "heading":
                    lower = line.strip().lower()
                    if any(kw in lower for kw in cls.SUMMARY_KW):
                        in_summary = True
                    else:
                        if in_summary:
                            return items
                        in_summary = False
                    continue
                if in_summary and cls_label == "summary_bullet":
                    bullet_text = line.strip().lstrip("•").strip()
                    items.append({
                        "para_idx": rec["para_idx"],
                        "line_idx": i,
                        "text": bullet_text,
                    })
        return items

    @classmethod
    def extract_skills(cls, parsed: List[Dict]) -> List[Dict]:
        """Extract skill lines."""
        items = []
        in_skills = False
        for rec in parsed:
            for i, (line, cls_label) in enumerate(zip(rec["lines"], rec["classifications"])):
                if cls_label == "heading":
                    lower = line.strip().lower()
                    if any(kw in lower for kw in cls.SKILLS_KW):
                        in_skills = True
                    else:
                        if in_skills:
                            return items
                        in_skills = False
                    continue
                if in_skills and cls_label == "skill_line" and line.strip():
                    items.append({
                        "para_idx": rec["para_idx"],
                        "line_idx": i,
                        "text": line.strip(),
                    })
        return items

    @classmethod
    def extract_header_points(cls, parsed: List[Dict]) -> List[Dict]:
        """
        Extract the keyword strip below the contact line and before the first section heading.
        """
        items = []
        for rec in parsed:
            para_idx = rec["para_idx"]
            if para_idx == 0:
                continue
            for i, (line, cls_label) in enumerate(zip(rec["lines"], rec["classifications"])):
                stripped = line.strip()
                if not stripped:
                    continue
                if cls_label == "heading":
                    return items
                if "•" in stripped and cls_label == "text":
                    items.append({
                        "para_idx": para_idx,
                        "line_idx": i,
                        "text": stripped,
                    })
                    return items
        return items

    @staticmethod
    def full_text(doc: Document) -> str:
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    @classmethod
    def line_text_map(cls, doc: Document) -> Dict[Tuple[int, int], str]:
        """Return the actual visual line text for each paragraph/line in the document."""
        parsed = cls.parse_paragraphs(doc)
        line_map: Dict[Tuple[int, int], str] = {}
        for rec in parsed:
            for idx, line in enumerate(rec["lines"]):
                line_map[(rec["para_idx"], idx)] = line
        return line_map


# ======================================================================
# AI ENGINE
# ======================================================================

class AIEngine:
    def __init__(self, api_key: str, model: str = DEFAULT_MODEL):
        self.client = OpenAI(api_key=api_key)
        self.model = model

    def extract_jd_focus(self, job_description: str) -> str:
        prompt = f"""You are analyzing a job description for a resume rewrite system.

JOB DESCRIPTION:
{job_description}

Return ONLY valid JSON with this shape:
{{
  "target_title": "<short role label>",
  "must_have_skills": ["skill1", "skill2"],
  "preferred_skills": ["skill1", "skill2"],
  "domain_context": ["domain1", "domain2"],
  "business_outcomes": ["outcome1", "outcome2"],
  "quantification_signals": ["metric pattern or scale signal"],
  "rewrite_guardrails": ["constraint1", "constraint2"]
}}

Rules:
- Keep every list short and high-signal.
- Include only information explicitly supported by the JD.
- Do not infer employer-specific details not present in the JD.
"""
        try:
            raw = self._call_text(prompt, temperature=0.1)
            cleaned = re.sub(r"```json?\s*", "", raw)
            cleaned = re.sub(r"```\s*", "", cleaned)
            parsed = json.loads(cleaned)
            return json.dumps(parsed, ensure_ascii=True, indent=2)
        except Exception:
            return json.dumps({
                "target_title": "",
                "must_have_skills": [],
                "preferred_skills": [],
                "domain_context": [],
                "business_outcomes": [],
                "quantification_signals": [],
                "rewrite_guardrails": [],
            }, ensure_ascii=True, indent=2)

    def rewrite_bullets(self, role_key: str, bullets: List[str],
                        job_description: str,
                        aggressive_rewrite: bool = False) -> List[str]:
        min_quantified = len(bullets) if aggressive_rewrite else max(1, math.ceil(len(bullets) * 0.5))
        jd_focus = self.extract_jd_focus(job_description)
        numbered = "\n".join(f"{i+1}. {b}" for i, b in enumerate(bullets))
        rewrite_mode = (
            "You may substantially rewrite bullets to maximize alignment with the JD, but every rewrite must remain traceable to the original bullet and plausible for "
            f'"{role_key}". Prioritize JD-required skills, responsibilities, and measurable outcomes without changing the underlying experience.'
            if aggressive_rewrite else
            "Stay GROUNDED in the original resume content. Preserve the original domain, core technology stack, and responsibility unless the JD clearly supports a tighter rewording."
        )
        ai_guardrail = (
            "You may introduce JD-aligned AI / GenAI / platform language when plausible for the role context."
            if aggressive_rewrite else
            "Do NOT replace data engineering bullets with generic AI/LLM bullets unless the ORIGINAL bullet or ROLE / CONTEXT already indicates GenAI/LLM/ML work."
        )
        noun_guardrail = (
            "Preserve at least one original anchor per bullet whenever available: a concrete noun, platform, technology, domain, stakeholder, deliverable, or business process from the original bullet."
            if aggressive_rewrite else
            "Each rewritten bullet should retain at least one concrete noun, technology, or domain signal from the original bullet when available."
        )
        quality_bar = (
            "- Good: JD-first bullet with concrete technologies, responsibilities, and quantifiable impact that still reads like the same accomplishment as the original bullet."
            if aggressive_rewrite else
            "- Good: original accomplishment preserved, JD-relevant keywords added, measurable outcome included."
        )
        prompt = self._build_bullet_prompt(
            role_key=role_key,
            bullets=bullets,
            numbered=numbered,
            job_description=job_description,
            jd_focus=jd_focus,
            min_quantified=min_quantified,
            rewrite_mode=rewrite_mode,
            ai_guardrail=ai_guardrail,
            noun_guardrail=noun_guardrail,
            quality_bar=quality_bar,
        )
        rewritten = self._call_json_list(prompt, expected_len=len(bullets), temperature=0.2)

        if not self._bullet_batch_is_strong_enough(bullets, rewritten, job_description, min_quantified):
            retry_prompt = prompt + f"""

CORRECTION PASS:
- Several bullets were not grounded enough to the original resume, were missing metrics, or were not aligned enough to the JD.
- Rewrite the full set again.
- For EVERY bullet, preserve at least one original anchor and add 1-2 JD terms only when supported by the original bullet or role context.
- Ensure at least {min_quantified} bullets contain a concrete metric or scale signal.
- Avoid generic filler; each bullet should sound distinct and evidence-based.
"""
            rewritten = self._call_json_list(retry_prompt, expected_len=len(bullets), temperature=0.1)
        return rewritten

        min_quantified = len(bullets) if aggressive_rewrite else max(1, math.ceil(len(bullets) * 0.5))
        jd_focus = self.extract_jd_focus(job_description)
        numbered = "\n".join(f"{i+1}. {b}" for i, b in enumerate(bullets))
        rewrite_mode = (
            "You MAY rewrite bullets completely to match the JD as long as they remain plausible for "
            f'"{role_key}". Prioritize JD-required skills, responsibilities, and measurable outcomes over original wording.'
            if aggressive_rewrite else
            "Stay GROUNDED in the original resume content. Preserve the original domain, core technology stack, and responsibility unless the JD clearly supports a tighter rewording."
        )
        ai_guardrail = (
            "You may introduce JD-aligned AI / GenAI / platform language when plausible for the role context."
            if aggressive_rewrite else
            "Do NOT replace data engineering bullets with generic AI/LLM bullets unless the ORIGINAL bullet or ROLE / CONTEXT already indicates GenAI/LLM/ML work."
        )
        noun_guardrail = (
            "Retain original signals when useful, but you do not need to preserve original nouns or technologies if a stronger JD-aligned rewrite is more effective and still plausible."
            if aggressive_rewrite else
            "Each rewritten bullet should retain at least one concrete noun, technology, or domain signal from the original bullet when available."
        )
        quality_bar = (
            "- Good: JD-first bullet with concrete technologies, responsibilities, and quantifiable impact that fits the candidate's seniority and client context."
            if aggressive_rewrite else
            "- Good: original accomplishment preserved, JD-relevant keywords added, measurable outcome included."
        )
        prompt = f"""You are an expert resume writer and ATS optimisation specialist.

ROLE / CONTEXT:
{role_key}

ORIGINAL BULLET POINTS:
{numbered}

TARGET JOB DESCRIPTION:
{job_description}

JD FOCUS SUMMARY:
{jd_focus}

STRICT RULES:
1. Return EXACTLY {len(bullets)} bullet points — same count as input.
2. {rewrite_mode}
3. {ai_guardrail}
4. For each bullet:
   • RELEVANT → refine with stronger action verbs and clearer business impact.
   • PARTIALLY RELEVANT → align wording more strongly to the JD.
   • WEAK / GENERIC → improve specificity using plausible details from the role context or JD focus summary.
5. Prefer JD keyword overlap, but never at the cost of factual drift.
6. {noun_guardrail}
7. Wrap JD-matched keywords in **double asterisks** for bold.
   Example: Designed scalable **ETL pipelines** using **Python** and **AWS Glue**.
8. At least {min_quantified} of the {len(bullets)} bullets must include a concrete metric,
   percentage, count, scale, latency, SLA, time reduction, volume, or dollar impact.
   Reuse existing metrics whenever present. If a bullet has no explicit metric, add only a
   conservative and plausible metric grounded in the original bullet, role, or JD.
9. Keep bullets outcome-oriented. Tie the work to reliability, latency, scale, automation,
   reporting, customer impact, compliance, or cost when supported.
10. Start each bullet with a strong action verb; 1–2 lines.
11. Do NOT include the "•" bullet marker — just the text.
12. Output ONLY a JSON array of {len(bullets)} strings.

QUALITY BAR:
- Bad: generic buzzwords, invented AI work, or bullets that could fit any candidate.
{quality_bar}

OUTPUT:
"""
        return self._call_json_list(prompt, expected_len=len(bullets), temperature=0.2)

    def _build_bullet_prompt(self, role_key: str, bullets: List[str], numbered: str,
                             job_description: str, jd_focus: str, min_quantified: int,
                             rewrite_mode: str, ai_guardrail: str,
                             noun_guardrail: str, quality_bar: str) -> str:
        return f"""You are an expert resume writer and ATS optimisation specialist.

ROLE / CONTEXT:
{role_key}

ORIGINAL BULLET POINTS:
{numbered}

TARGET JOB DESCRIPTION:
{job_description}

JD FOCUS SUMMARY:
{jd_focus}

STRICT RULES:
1. Return EXACTLY {len(bullets)} bullet points - same count as input.
2. {rewrite_mode}
3. {ai_guardrail}
4. For each bullet:
   - RELEVANT -> refine with stronger action verbs and clearer business impact.
   - PARTIALLY RELEVANT -> align wording more strongly to the JD.
   - WEAK / GENERIC -> improve specificity using plausible details from the role context or JD focus summary.
5. Prefer JD keyword overlap, but never at the cost of factual drift.
6. {noun_guardrail}
7. Maximum JD alignment matters: each bullet should include the strongest JD-matching skills, platforms, responsibilities, or outcomes that are genuinely supported by the original bullet or role context.
8. Do not fabricate tools, ownership, scope, seniority, compliance exposure, customer context, or AI/ML work that is not reasonably supported by the original bullet, role context, or JD focus summary.
9. Treat the original bullet as evidence. Preserve the accomplishment's underlying action and context even when the wording changes significantly.
10. If the original bullet already contains a metric, preserve it unless a tighter wording is needed. If it does not, add only a conservative metric or scale signal that is plausible from the original bullet, role context, or explicit JD scale language.
11. Prefer quantified outcomes tied to throughput, latency, uptime, automation, delivery speed, data volume, adoption, revenue, cost, compliance, customer impact, or team scope when supported.
12. Wrap JD-matched keywords in **double asterisks** for bold.
   Example: Designed scalable **ETL pipelines** using **Python** and **AWS Glue**.
13. At least {min_quantified} of the {len(bullets)} bullets must include a concrete metric, percentage, count, scale, latency, SLA, time reduction, volume, or dollar impact.
14. Keep bullets outcome-oriented. Tie the work to reliability, latency, scale, automation, reporting, customer impact, compliance, or cost when supported.
15. Start each bullet with a strong action verb; 1-2 lines.
16. Do NOT include the bullet marker - just the text.
17. Output ONLY a JSON array of {len(bullets)} strings.

ADDITIONAL GROUNDING CHECK:
- Before finalising each bullet, ask: "Can I point to the original bullet evidence for this claim?"
- If not, remove or soften that claim.
- A strong rewrite should still feel like an upgraded version of the original accomplishment, not a brand-new project.

QUALITY BAR:
- Bad: generic buzzwords, invented AI work, or bullets that could fit any candidate.
{quality_bar}

OUTPUT:
"""

    def _extract_keywords(self, text: str) -> set:
        return {
            token for token in re.findall(r"[A-Za-z][A-Za-z0-9+#./-]{2,}", text.lower())
            if token not in STOPWORDS
        }

    def _has_metric(self, text: str) -> bool:
        return bool(METRIC_PATTERN.search(text))

    def _bullet_batch_is_strong_enough(self, originals: List[str], rewritten: List[str],
                                       job_description: str, min_quantified: int) -> bool:
        jd_terms = self._extract_keywords(job_description)
        quantified = 0
        grounded = 0
        jd_aligned = 0

        for original, candidate in zip(originals, rewritten):
            orig_terms = self._extract_keywords(original)
            cand_terms = self._extract_keywords(candidate)

            if self._has_metric(candidate):
                quantified += 1
            if orig_terms and cand_terms.intersection(orig_terms):
                grounded += 1
            if jd_terms and cand_terms.intersection(jd_terms):
                jd_aligned += 1

        return (
            quantified >= min_quantified and
            grounded >= max(1, math.ceil(len(originals) * 0.8)) and
            jd_aligned >= max(1, math.ceil(len(originals) * 0.8))
        )

    def rewrite_summary(self, summary_bullets: List[str],
                        job_description: str,
                        aggressive_rewrite: bool = False) -> List[str]:
        """Rewrite professional summary bullets to align with the JD."""
        min_quantified = len(summary_bullets) if aggressive_rewrite else max(1, math.ceil(len(summary_bullets) * 0.3))
        jd_focus = self.extract_jd_focus(job_description)
        canonical_title = "Senior Data and Gen AI Engineer"
        numbered = "\n".join(f"{i+1}. {b}" for i, b in enumerate(summary_bullets))
        summary_mode = (
            "You MAY rewrite the summary bullets completely to position the candidate for the JD. Focus on the target role, required platforms, and measurable outcomes."
            if aggressive_rewrite else
            "Rewrite each bullet to align with the job description while preserving the candidate's core experience and seniority level."
        )
        summary_grounding = (
            "Do not stay tied to the original summary wording. Use the JD and overall role history to produce a sharper positioning statement."
            if aggressive_rewrite else
            "Stay grounded in the original resume summary. Do NOT turn a data engineer into an AI engineer unless the original summary or role history clearly supports that positioning."
        )
        summary_signal = (
            "You may emphasize JD-critical capabilities even if they were not mentioned in the original summary, as long as they are plausible from the candidate's experience."
            if aggressive_rewrite else
            "Each bullet should preserve at least one original capability, platform, or domain signal."
        )
        prompt = f"""You are an expert resume writer and ATS optimisation specialist.

SECTION: Professional Summary

ORIGINAL SUMMARY BULLET POINTS:
{numbered}

TARGET JOB DESCRIPTION:
{job_description}

JD FOCUS SUMMARY:
{jd_focus}

STRICT RULES:
1. Return EXACTLY {len(summary_bullets)} summary bullet points — same count as input.
2. {summary_mode}
3. {summary_grounding}
4. The FIRST summary bullet must preserve the candidate's professional identity as
   "{canonical_title}". Do NOT rename the candidate to "AI Engineer", "ML Engineer",
   or another title in that first bullet.
5. Emphasise skills, technologies, and achievements that match the JD.
6. At least {min_quantified} of the {len(summary_bullets)} summary bullets must contain a
   concrete metric, count, percentage, scale, or outcome where plausible.
7. {summary_signal}
8. Wrap JD-matched keywords in **double asterisks** for bold.
9. Keep each bullet concise and professional (1–2 sentences).
10. Do NOT include the "•" marker — just the text.
11. Output ONLY a JSON array of {len(summary_bullets)} strings.

OUTPUT:
"""
        return self._call_json_list(prompt, expected_len=len(summary_bullets), temperature=0.2)

    def update_skills(self, current_skills_lines: List[str],
                      job_description: str) -> List[str]:
        current_skills_text = "\n".join(current_skills_lines)
        jd_focus = self.extract_jd_focus(job_description)
        prompt = f"""You are an expert resume writer.

CURRENT SKILLS SECTION (each line is a category like "Category: skill1, skill2, ..."):
{current_skills_text}

JOB DESCRIPTION:
{job_description}

JD FOCUS SUMMARY:
{jd_focus}

INSTRUCTIONS:
1. You MAY completely rewrite the professional skills section to align with the JD.
2. Extract and prioritise the most relevant technologies, platforms, domains, and methods from the JD.
3. Remove low-value or irrelevant skills, and add high-value JD-aligned skills that fit the candidate profile.
4. IMPORTANT: Keep the EXACT SAME number of lines as the original to preserve the document layout.
5. You MAY rename category labels completely. Each line must still follow the format:
   "Category Label: skill1, skill2, skill3"
6. Make the section feel curated for the target role, not like a generic inventory dump.
7. Front-load the strongest JD-matching skills in each line.
8. Wrap JD-matched keywords in **double asterisks**.
9. Return plain text only. Do NOT use code fences, markdown labels, or ```plaintext.
10. Return ONLY the updated skills text — no extra commentary.
"""
        raw = self._call_text(prompt)
        lines = _clean_skills_output(raw)
        if len(lines) > len(current_skills_lines):
            lines = lines[:len(current_skills_lines)]
        while len(lines) < len(current_skills_lines):
            fallback = current_skills_lines[len(lines)]
            lines.append(fallback)
        return lines

    def rewrite_header_points(self, current_header_line: str,
                              job_description: str) -> str:
        jd_focus = self.extract_jd_focus(job_description)
        current_points = [part.strip() for part in current_header_line.split("•") if part.strip()]
        target_count = max(8, len(current_points)) if current_points else 10
        prompt = f"""You are an expert resume writer.

CURRENT HEADER KEYWORD STRIP:
{current_header_line}

TARGET JOB DESCRIPTION:
{job_description}

JD FOCUS SUMMARY:
{jd_focus}

INSTRUCTIONS:
1. Rewrite the header keyword strip completely to align with the JD.
2. Return a single line containing exactly {target_count} concise keyword clusters separated by " • ".
3. Each cluster should be 1-4 words and reflect high-value JD keywords, platforms, or capabilities.
4. Prioritise the strongest match terms first.
5. Avoid duplicates, filler, or generic soft skills.
6. Do NOT use code fences or extra commentary.
7. Return ONLY the single rewritten line.
"""
        raw = self._call_text(prompt, temperature=0.2).strip()
        cleaned = re.sub(r"^```[A-Za-z0-9_-]*\s*", "", raw)
        cleaned = re.sub(r"\s*```$", "", cleaned).strip()
        parts = [part.strip() for part in cleaned.split("•") if part.strip()]
        if len(parts) > target_count:
            parts = parts[:target_count]
        while len(parts) < target_count:
            fallback_idx = len(parts) % max(len(current_points), 1)
            fallback = current_points[fallback_idx] if current_points else "Platform Engineering"
            parts.append(fallback)
        return " • ".join(parts)

    def ats_analysis(self, resume_text: str, job_description: str) -> Dict:
        prompt = f"""You are an ATS analysis engine.

RESUME:
{resume_text}

JOB DESCRIPTION:
{job_description}

Return JSON:
{{
  "ats_score": <int 0-100>,
  "missing_skills": ["skill1", ...],
  "keyword_freq": {{"keyword": <count_in_resume>, ...}}
}}
Output ONLY valid JSON.
"""
        raw = self._call_text(prompt)
        try:
            cleaned = re.sub(r"```json?\s*", "", raw)
            cleaned = re.sub(r"```\s*", "", cleaned)
            return json.loads(cleaned)
        except json.JSONDecodeError:
            return {"ats_score": 0, "missing_skills": [], "keyword_freq": {}}

    def generate_cover_letter(self, resume_text: str,
                              job_description: str) -> str:
        prompt = f"""You are a professional cover letter writer.

RESUME:
{resume_text}

JOB DESCRIPTION:
{job_description}

Write a tailored cover letter (3–4 paragraphs):
1. Open with enthusiasm for the role.
2. Highlight 2–3 key resume achievements matching the JD.
3. Demonstrate knowledge of the company/role.
4. Close with a confident call to action.
Professional tone. Output ONLY the cover letter.
"""
        return self._call_text(prompt)

    # ── helpers ───────────────────────────────────────────────────────
    def _call_text(self, prompt: str, temperature: float = 0.4) -> str:
        resp = self.client.chat.completions.create(
            model=self.model,
            messages=[{"role": "user", "content": prompt}],
            temperature=temperature,
        )
        return resp.choices[0].message.content.strip()

    def _call_json_list(self, prompt: str, expected_len: int, temperature: float = 0.4) -> List[str]:
        raw = self._call_text(prompt, temperature=temperature)
        try:
            cleaned = re.sub(r"```json?\s*", "", raw)
            cleaned = re.sub(r"```\s*", "", cleaned)
            parsed = json.loads(cleaned)
            if isinstance(parsed, list):
                if len(parsed) > expected_len:
                    parsed = parsed[:expected_len]
                while len(parsed) < expected_len:
                    parsed.append(parsed[-1] if parsed else "")
                return [str(x) for x in parsed]
        except json.JSONDecodeError:
            pass
        lines = [l.strip().lstrip("0123456789.-) ") for l in raw.split("\n") if l.strip()]
        while len(lines) < expected_len:
            lines.append(lines[-1] if lines else "")
        return lines[:expected_len]


# ======================================================================
# DOCUMENT WRITER
# ======================================================================

class DocumentWriter:
    @staticmethod
    def sync_paragraph_layout(source_doc: Document, target_doc: Document):
        """Copy paragraph properties from the original document after rebuilds."""
        for src_para, dst_para in zip(source_doc.paragraphs, target_doc.paragraphs):
            src_p = src_para._element
            dst_p = dst_para._element
            src_pPr = src_p.find(qn("w:pPr"))
            dst_pPr = dst_p.find(qn("w:pPr"))

            if dst_pPr is not None:
                dst_p.remove(dst_pPr)
            if src_pPr is not None:
                dst_p.insert(0, copy.deepcopy(src_pPr))

    @staticmethod
    def verify_updates(doc: Document,
                       updates: Dict[Tuple[int, int], Tuple[str, bool]]) -> Dict[Tuple[int, int], str]:
        """
        Return the actual applied line text for each updated line.

        If a line cannot be found exactly as intended, we still return the
        document's real line text so the UI preview matches the downloaded file.
        """
        actual_lines = DocumentParser.line_text_map(doc)
        verified: Dict[Tuple[int, int], str] = {}

        for key, (expected_text, is_bullet) in updates.items():
            actual_text = actual_lines.get(key, "")
            verified[key] = actual_text

            expected_norm = _normalize_line_text(expected_text, is_bullet)
            actual_norm = _normalize_line_text(actual_text, is_bullet)

            if expected_norm == actual_norm:
                continue

            para_idx, line_idx = key
            nearby_keys = [
                (para_idx, line_idx - 1),
                (para_idx, line_idx + 1),
                (para_idx, line_idx - 2),
                (para_idx, line_idx + 2),
            ]
            for nearby in nearby_keys:
                candidate = actual_lines.get(nearby)
                if not candidate:
                    continue
                candidate_norm = _normalize_line_text(candidate, is_bullet)
                if candidate_norm == expected_norm:
                    verified[key] = candidate
                    break

        return verified

    @staticmethod
    def save_to_bytes(doc: Document) -> bytes:
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()


# ======================================================================
# STREAMLIT UI
# ======================================================================

def main():
    st.set_page_config(page_title="Resume Aligner", page_icon="📄", layout="wide")
    st.title("📄 Resume Aligner")
    st.caption(
        "Upload your resume and paste a job description — "
        "the AI aligns your resume while preserving its exact structure."
    )

    # ── Sidebar ──────────────────────────────────────────────────────
    with st.sidebar:
        st.header("⚙️ Configuration")
        api_key = st.text_input(
            "OpenAI API Key", type="password",
            value=os.getenv("OPENAI_API_KEY", ""),
            help="Set in .env or paste here.",
        )
        model = st.selectbox(
            "Model",
            ["gpt-4o", "gpt-4o-mini", "gpt-4-turbo", "gpt-3.5-turbo"],
            index=0,
        )
        st.divider()
        st.markdown(
            "**How it works**\n"
            "1. Upload `.docx` resume\n"
            "2. Paste job description\n"
            "3. Click **Align Resume**\n"
            "4. Download optimised file"
        )

    # ── Inputs ───────────────────────────────────────────────────────
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Resume Upload")
        uploaded = st.file_uploader("Upload resume (.docx)", type=["docx"])
    with c2:
        st.subheader("Job Description")
        jd = st.text_area("Paste job description", height=300,
                          placeholder="Copy and paste the full job description...")

    if not api_key:
        st.info("Enter your OpenAI API key in the sidebar.")
        return
    if not uploaded or not jd.strip():
        st.info("Upload a resume and paste a job description to continue.")
        return

    # ── GO ───────────────────────────────────────────────────────────
    if st.button("🚀 Align Resume", type="primary", use_container_width=True):
        file_bytes = uploaded.read()
        try:
            ai = AIEngine(api_key=api_key, model=model)
        except Exception as e:
            st.error(f"OpenAI init failed: {e}")
            return

        # ── Step 1: Parse ────────────────────────────────────────────
        with st.status("Parsing resume...", expanded=True) as status:
            doc = DocumentParser.load_document(file_bytes)
            parsed = DocumentParser.parse_paragraphs(doc)
            roles = DocumentParser.extract_roles(parsed)
            summary_bullets = DocumentParser.extract_summary_bullets(parsed)
            skills = DocumentParser.extract_skills(parsed)
            header_points = DocumentParser.extract_header_points(parsed)
            total_b = sum(len(r["bullets"]) for r in roles.values())
            st.write(f"**{len(roles)} roles**, **{total_b} experience bullets**, "
                     f"**{len(summary_bullets)} summary bullets**, **{len(skills)} skill lines**")
            for rk, rv in roles.items():
                st.write(f"  • *{rk[:80]}* — {len(rv['bullets'])} bullets")
            status.update(label="Parsed", state="complete")

        # We will collect all updates as (para_idx, line_idx) → new_text
        # and whether each is a bullet or not
        all_updates: Dict[Tuple[int,int], Tuple[str, bool]] = {}
        original_map: Dict[Tuple[int,int], str] = {}
        applied_map: Dict[Tuple[int,int], str] = {}

        with st.status("Refreshing header keywords...", expanded=True) as status:
            if header_points:
                for hinfo in header_points:
                    key = (hinfo["para_idx"], hinfo["line_idx"])
                    try:
                        rewritten_header = ai.rewrite_header_points(hinfo["text"], jd)
                    except Exception as e:
                        st.warning(f"Header strip error: {e}")
                        rewritten_header = hinfo["text"]
                    all_updates[key] = (rewritten_header, False)
                    original_map[key] = hinfo["text"]
                st.write(f"Updated {len(header_points)} header keyword line(s).")
            else:
                st.write("No header keyword strip found.")
            status.update(label="Header done", state="complete")

        # ── Step 2: Rewrite summary ──────────────────────────────────
        with st.status("Aligning Professional Summary...", expanded=True) as status:
            if summary_bullets:
                originals = [b["text"] for b in summary_bullets]
                try:
                    rewritten = ai.rewrite_summary(originals, jd, aggressive_rewrite=True)
                except Exception as e:
                    st.warning(f"Summary error: {e}")
                    rewritten = originals
                for binfo, new_t in zip(summary_bullets, rewritten):
                    key = (binfo["para_idx"], binfo["line_idx"])
                    all_updates[key] = (new_t, True)  # is_bullet=True (has •)
                    original_map[key] = binfo["text"]
                st.write(f"Rewrote {len(summary_bullets)} summary bullets.")
            else:
                st.write("No summary bullets found.")
            status.update(label="Summary done", state="complete")

        # ── Step 3: Rewrite experience bullets ───────────────────────
        with st.status("Aligning experience bullets...", expanded=True) as status:
            prog = st.progress(0)
            rlist = list(roles.items())
            for ri, (rk, rv) in enumerate(rlist):
                if not rv["bullets"]:
                    prog.progress((ri+1)/max(len(rlist),1))
                    continue
                st.write(f"Processing: *{rk[:80]}*…")
                originals = [b["text"] for b in rv["bullets"]]
                aggressive_rewrite = ri < 2
                try:
                    rewritten = ai.rewrite_bullets(rk, originals, jd, aggressive_rewrite=aggressive_rewrite)
                except Exception as e:
                    st.warning(f"Error for '{rk[:40]}': {e}")
                    rewritten = originals
                for binfo, new_t in zip(rv["bullets"], rewritten):
                    key = (binfo["para_idx"], binfo["line_idx"])
                    all_updates[key] = (new_t, True)
                    original_map[key] = binfo["text"]
                prog.progress((ri+1)/max(len(rlist),1))
            status.update(label=f"Aligned {total_b} experience bullets", state="complete")

        # ── Step 4: Skills ───────────────────────────────────────────
        with st.status("Updating skills...", expanded=True) as status:
            if skills:
                current_skill_lines = [s["text"] for s in skills]
                try:
                    new_lines = ai.update_skills(current_skill_lines, jd)
                    for si, new_l in zip(skills, new_lines):
                        key = (si["para_idx"], si["line_idx"])
                        all_updates[key] = (new_l, False)  # is_bullet=False
                    st.write(f"Updated {min(len(new_lines), len(skills))} skill lines.")
                except Exception as e:
                    st.warning(f"Skills error: {e}")
            else:
                st.write("No skills section found.")
            status.update(label="Skills done", state="complete")

        # ── Step 5: Rebuild document ─────────────────────────────────
        with st.status("Rebuilding document (preserving formatting)…", expanded=True) as status:
            original_doc = DocumentParser.load_document(file_bytes)
            aligned_doc = DocumentParser.load_document(file_bytes)

            # Group updates by paragraph
            updates_by_para: Dict[int, Dict[int, Tuple[str, bool]]] = {}
            for (pidx, lidx), (txt, is_bul) in all_updates.items():
                updates_by_para.setdefault(pidx, {})[lidx] = (txt, is_bul)

            for pidx in sorted(updates_by_para.keys()):
                para = aligned_doc.paragraphs[pidx]
                line_map = updates_by_para[pidx]

                # Determine if this paragraph's changes are bullets or skill lines
                # (mixed is handled: use per-line is_bullet flag)
                # ParagraphRebuilder needs a simple {line_idx: text} map + is_bullet
                # Since is_bullet can vary per line, we call rebuild once per line type group
                bullet_replacements = {l: t for l, (t, ib) in line_map.items() if ib}
                skill_replacements = {l: t for l, (t, ib) in line_map.items() if not ib}

                if bullet_replacements and skill_replacements:
                    # Both in same paragraph (unlikely but handle it):
                    # Do bullet first, then skill
                    ParagraphRebuilder.rebuild(para, bullet_replacements, is_bullet=True)
                    # Re-parse for skill pass
                    ParagraphRebuilder.rebuild(para, skill_replacements, is_bullet=False, is_skill=True)
                elif bullet_replacements:
                    ParagraphRebuilder.rebuild(para, bullet_replacements, is_bullet=True)
                elif skill_replacements:
                    ParagraphRebuilder.rebuild(para, skill_replacements, is_bullet=False, is_skill=True)

            DocumentWriter.sync_paragraph_layout(original_doc, aligned_doc)
            applied_map = DocumentWriter.verify_updates(aligned_doc, all_updates)
            doc_bytes = DocumentWriter.save_to_bytes(aligned_doc)
            status.update(label="Document rebuilt — formatting preserved", state="complete")

        # ── Step 6: ATS ──────────────────────────────────────────────
        with st.status("ATS analysis…", expanded=True) as status:
            a_text = DocumentParser.full_text(DocumentParser.load_document(doc_bytes))
            try:
                ats = ai.ats_analysis(a_text, jd)
            except Exception as e:
                ats = {"ats_score": 0, "missing_skills": [], "keyword_freq": {}}
            status.update(label="ATS done", state="complete")

        # ── Step 7: Cover letter ─────────────────────────────────────
        with st.status("Cover letter…", expanded=True) as status:
            try:
                cover = ai.generate_cover_letter(a_text, jd)
            except Exception as e:
                cover = f"Error: {e}"
            status.update(label="Cover letter done", state="complete")

        # ══════════════════════════════════════════════════════════════
        # RESULTS
        # ══════════════════════════════════════════════════════════════
        st.divider()
        st.header("Results")

        st.download_button(
            "⬇️ Download Aligned Resume (.docx)", data=doc_bytes,
            file_name="aligned_resume.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary", use_container_width=True,
        )

        # ATS
        st.subheader("ATS Score")
        score = ats.get("ats_score", 0)
        color = "🟢" if score >= 75 else ("🟡" if score >= 50 else "🔴")
        st.metric("Match Score", f"{score} / 100", delta=color)

        missing = ats.get("missing_skills", [])
        if missing:
            st.subheader("Missing Skills")
            st.warning("These JD skills are not in your resume:")
            for s in missing:
                st.markdown(f"- {s}")
        else:
            st.success("No critical missing skills detected.")

        kw = ats.get("keyword_freq", {})
        if kw:
            st.subheader("Keyword Frequency")
            df = pd.DataFrame([{"Keyword": k, "Count": v} for k, v in kw.items()])
            df = df.sort_values("Count", ascending=False).reset_index(drop=True)
            st.dataframe(df, use_container_width=True, hide_index=True)

        st.subheader("Tailored Cover Letter")
        st.text_area("Cover Letter", value=cover, height=350)

        # Diff
        with st.expander("📋 Bullet-by-Bullet Changes"):
            # Summary
            if summary_bullets:
                st.markdown("**Professional Summary**")
                for binfo in summary_bullets:
                    key = (binfo["para_idx"], binfo["line_idx"])
                    orig = original_map.get(key, binfo["text"])
                    fallback_text, is_bullet = all_updates.get(key, (orig, True))
                    shown_text = applied_map.get(key, fallback_text)
                    clean_new = _normalize_line_text(shown_text, is_bullet)
                    if _normalize_line_text(orig, True) != clean_new:
                        st.markdown(f"~~{orig[:200]}~~")
                        st.markdown(f"→ {shown_text[:200]}")
                    else:
                        st.markdown(f"✓ {orig[:200]} *(unchanged)*")
                st.divider()

            # Experience
            for rk, rv in roles.items():
                st.markdown(f"**{rk[:120]}**")
                for binfo in rv["bullets"]:
                    key = (binfo["para_idx"], binfo["line_idx"])
                    orig = original_map.get(key, binfo["text"])
                    fallback_text, is_bullet = all_updates.get(key, (orig, True))
                    shown_text = applied_map.get(key, fallback_text)
                    clean_new = _normalize_line_text(shown_text, is_bullet)
                    if _normalize_line_text(orig, True) != clean_new:
                        st.markdown(f"~~{orig[:200]}~~")
                        st.markdown(f"→ {shown_text[:200]}")
                    else:
                        st.markdown(f"✓ {orig[:200]} *(unchanged)*")
                st.divider()


if __name__ == "__main__":
    from web_app import main as web_main
    web_main()
